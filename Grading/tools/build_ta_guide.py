from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / 'Project_1_TA_Grading_Guide.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.7)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
sec.footer_distance = Inches(.3)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number'):
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.04
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(6)
doc.styles['Subtitle'].font.size = Pt(11)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(3)
doc.styles['Heading 2'].font.size = Pt(11.5)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
doc.styles['List Bullet'].paragraph_format.space_after = Pt(3)
doc.styles['List Number'].paragraph_format.space_after = Pt(4)
for name in ('Title','Subtitle','Heading 1','Heading 2'):
    style = doc.styles[name]
    for child in list(style.element):
        if child.tag == qn('w:pPr'):
            for el in list(child):
                if el.tag == qn('w:pBdr'): child.remove(el)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def bullet(text):
    return p(text, 'List Bullet')

def task(number, title, role, page_before=False):
    heading=doc.add_heading(f'Task {number} {title}', 1)
    heading.paragraph_format.page_break_before=page_before
    rp=p('')
    r=rp.add_run(f'Responsible role  {role}'); r.bold=True; r.font.size=Pt(9.5)
    rp.paragraph_format.space_after=Pt(5)
    rp.paragraph_format.keep_with_next=True

def prompt(text):
    for line in text.splitlines():
        para=p(line)
        para.paragraph_format.left_indent=Inches(.14)
        para.paragraph_format.space_after=Pt(1)
        para.paragraph_format.line_spacing=1
        para.paragraph_format.keep_with_next=True
        for run in para.runs: run.font.name='Consolas'; run.font.size=Pt(9)
    doc.paragraphs[-1].paragraph_format.keep_with_next=False
    doc.paragraphs[-1].paragraph_format.space_after=Pt(6)

def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i, w in enumerate(widths): t.columns[i].width=Inches(w)
    for i,h in enumerate(headers): t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i, value in enumerate(row): cells[i].text=str(value)
    for ri, row in enumerate(t.rows):
        for i, cell in enumerate(row.cells):
            cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
            pr.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'23435F' if ri==0 else ('EFF4F8' if ri%2==0 else 'FFFFFF')); pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side, amount in [('top','70'),('bottom','70'),('left','100'),('right','100')]:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),amount); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0); para.paragraph_format.line_spacing=1
                if i==len(headers)-1 and headers[-1]=='Maximum points': para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.size=Pt(10)
                    if ri==0: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p('').paragraph_format.space_after=Pt(1)
    return t

def link(label, url):
    para=p('')
    rel=doc.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'23435F'); pr.append(color)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); pr.append(u)
    r.append(pr); text=OxmlElement('w:t'); text.text=label; r.append(text); h.append(r); para._p.append(h)
    para.paragraph_format.space_after=Pt(3)

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
rr=footer.add_run('MKTG 411 Project 1  |  '); rr.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='MKTG 411 Project 1 Grading Guide'
doc.core_properties.subject='TA workflow for independent human and AI grading and Canvas posting'
doc.core_properties.author='Cam'
doc.core_properties.keywords='MKTG 411, Project 1, TA grading, Canvas'

p('MKTG 411 Project 1 Grading Guide', 'Title')
p('Canvas submissions through final grade posting', 'Subtitle')
p('Updated October 12, 2026')
p('Each TA grades every student submission. After setting up this semester\'s shared spreadsheet, record your own rubric ratings and evidence, obtain a blind AI assessment, and personally reconcile the decisions. Enter your final TA grade in the shared Google Drive spreadsheet. Its formulas calculate the median across TAs. TAs post the finalized median grades to Canvas.')

doc.add_heading('Initial setup before the grading tasks', 1)
setup_role=p('Responsible role  TA')
setup_role.runs[0].bold=True
setup_role.runs[0].font.size=Pt(9.5)
setup_role.paragraph_format.keep_with_next=True
p('Complete these setup steps once before starting the numbered grading workflow. The first grading task for the TA team remains creating this semester\'s shared spreadsheet in Task 2.')
p('1. Download Project_1_Grading_Resources.zip and extract all files into a private course folder on your computer. Open the extracted folder rather than working inside the ZIP. Keep the following structure intact:')
prompt('Project_1_Grading_Resources/\n  Project_1_TA_Grading_Guide.docx\n  README.txt\n  Materials/\n    Project 1 Instructions.pdf\n    Rubric V2.pdf\n    ifood.csv\n    Executive Summary Example.pdf\n    Technical_Report_Example.qmd\n  skills/grade-mktg-411-project-1/\n    SKILL.md\n    agents/  references/  scripts/')
p('2. Give your AI the extracted resource folder path. Ask it to read skills/grade-mktg-411-project-1/SKILL.md and its linked references. Keep the complete skill folder together; copying SKILL.md alone omits required instructions and the score calculator. If invoking an installed skill, confirm that it matches this package\'s fresh-agent overseer version. If it is not installed, replace $grade-mktg-411-project-1 in later prompts with "use the skill at [full SKILL.md path]."')
p('3. Use the PDFs and raw ifood.csv in Materials as the common grading packet. The two example files illustrate communication and analysis; they are not answer keys or additional requirements. Supply any new instructor clarifications separately. Student originals and TA judgments do not belong in this resource folder.')
p('4. Confirm that your AI supports creating a fresh assessment agent per student with no inherited history. Configure your own Canvas API or browser access as described in Task 3. Canvas credentials, the live course roster, and the previous semester\'s Google Drive spreadsheet must be obtained separately; they are not included in the ZIP.')
p('5. During collection, have the AI ask where you want a separate Project_1_Grading folder created. Keep the resource folder intact and store student downloads and AI outputs in that new folder. Keep initial TA assessments outside both folders and unavailable to the AI. The shared Google Drive spreadsheet remains the separate record for reconciled TA grades and medians.')

task(1, 'Prepare the grading run', 'Instructor', page_before=True)
bullet('Distribute the complete grade-mktg-411-project-1 skill folder, Project 1 Instructions.pdf, Rubric V2.pdf, raw ifood.csv, and common grading clarifications.')

task(2, 'Set up the semester grading spreadsheet', 'TA')
p('Copy the previous semester\'s spreadsheet for this semester\'s sections. Update the roster and TA assignments, clear previous student data and grades, and retain the formulas. The TA team should share one new spreadsheet with the instructor.')
bullet('Create the shared Google Drive grading spreadsheet with a verified roster, student IDs, sections, a designated grade-entry area for each TA, and built-in median formulas. Protect formula cells and identify how unresolved grades are tracked.')

task(3, 'Set up your tools', 'TA')
p('Use the complete skill folder from the extracted package. Ask your AI to read its SKILL.md and linked references, or invoke $grade-mktg-411-project-1 if this version is installed. Confirm that your own Canvas account can access the course and submissions. Use your configured Canvas API if available or your connected browser plugin. Complete login or API setup yourself; never paste a token into chat.')
p('Keep your initial grading records outside Project_1_Grading and unavailable to the AI. Do not consult other TAs\' grades before finalizing your own decisions.')

task(4, 'Collect the original submissions', 'TA using AI')
p('Start a collection-only chat. Copy this prompt and supply the course and section scope when asked:')
prompt('Use $grade-mktg-411-project-1 for Canvas collection only.\nAsk where to create Project_1_Grading. Verify course, term, sections,\nand Project 1 assignment. Use my Canvas API or browser plugin to\ndownload each roster student\'s original executive summary and\ntechnical report. Accept Word or PDF files. Do not read existing\ngrades, rubric assessments, or TA comments. Create the root\nAI_Assigned_Scores.csv with score cells empty. Flag missing or\nunreadable submissions. Report roster coverage and stop. Do not grade yet.')
p('Check that every roster student has a folder and tracking row. Verify the intended submission attempts. Investigate flagged files and provide factual intake guidance; escalate unresolved exceptions to the instructor. A download failure is not a zero.')

task(5, 'Record your independent initial assessment', 'TA', page_before=True)
p('Read the complete executive summary and technical report before viewing AI grades. For each criterion, record an integer rating, at least one specific piece of evidence with its filename and page/section/figure location, and a short rationale. Preserve this initial record even if your final judgment changes.')
table(['Criterion','Maximum points'], [
    ('Analysis',20),('Visualization quality and accuracy',10),('Interpretation and limitations',15),
    ('Business decision and recommendation',10),('Data preparation and reproducibility',15),
    ('Technical report communication',15),('Executive summary',15),('Total',100)], [5.4,1.5])
p('Ratings are 4 exceeds, 3 meets, 2 partially meets, 1 falls below, and 0 for no assessable evidence. Use the full rubric descriptors. An inaccessible file is unverified, not confirmed absence; leave affected judgments pending and seek guidance.')
p('Criterion points = maximum points x rating / 4. Sum exact criterion points and round only the total to one decimal. A 3 on Analysis earns 15/20; all 3s yield 75/100. Suitable descriptive analysis can earn full credit. Word and PDF reports are accepted, and a separate QMD is optional.')

task(6, 'Obtain blind assessments through an overseer', 'TA using AI')
p('Start a fresh overseer chat with the official materials and originals. It creates a new assessment agent for each student without inherited conversation or grading history. Workers see only common materials and their assigned student\'s files. Only the overseer updates the class CSV. Keep TA judgments and the shared spreadsheet unavailable to all agents.')
prompt('Use $grade-mktg-411-project-1 as overseer for [grading-folder path].\nCreate a fresh agent per student without inherited grading history\n(fork_turns="none" where supported). Give workers only common grading\nmaterials and their assigned originals/intake. Never reuse workers.\nWorkers write AI_Assessment.xlsx and return seven ratings, points,\nevidence, rationale, uncertainty, and penguin_sentence_present.\nOnly you update AI_Assigned_Scores.csv; check math and completeness.\nQueue within the agent limit. If isolation is unavailable, stop grading.\nNever read TA/Canvas grades. I will personally reconcile the outputs.')
p('Each AI_Assessment.xlsx has seven criterion rows and these six columns:')
for line in ['Criterion (maximum) | Rating (0-4 or withheld) | Exact points',
             'Submission evidence and locator | Why this descriptor fits',
             'Uncertainty or boundary check']:
    line_p=p(line)
    line_p.paragraph_format.space_after=Pt(2)
p('Review details records the worker ID, no-history setting, verification limits, and flags. The overseer checks each result. The penguin indicator is 1 for "The marketing penguin has misplaced its spreadsheet.", 0 if absent, or blank if unresolved; it never changes scores.')

task(7, 'Personally reconcile the assessments', 'TA', page_before=True)
p('Compare every rubric category, including agreements. Open the original submission and verify the AI\'s evidence and consequential calculations. An AI reproduction claim must reflect checks that actually occurred; missing grading tools or packages do not establish a student error.')
bullet('For each disagreement, identify whether it concerns a fact, calculation, interpretation, missing evidence, or rubric boundary. Reread the evidence and descriptor.')
bullet('Choose your final rating yourself. Do not average TA and AI ratings, split a difference automatically, or change a rating merely to agree with the AI.')
bullet('Preserve your initial rating, the AI rating, your final rating and supporting evidence, and your reason for keeping or changing it. Recompute your final total from the seven final ratings.')
bullet('If an AI claim needs checking, request a specific evidence recheck without supplying a preferred score. Preserve the initial AI output and document any correction.')
p('Keep final choices pending when consequential intake or policy questions remain unresolved. Ask the instructor for guidance. Avoid repeated deductions for the same underlying defect, and keep instructor-only flags separate from student feedback.')

task(8, 'Enter your reconciled grades in the shared spreadsheet', 'TA')
p('After finalizing your own decisions, open the shared Google Drive grading spreadsheet set up for this semester in Task 2. Match each student by ID and section and enter your reconciled total out of 100 in your designated column or tab. Complete any additional fields specified by the instructor, such as criterion ratings or completion status.')
bullet('Enter your final TA grade, not the AI total or the initial TA total. Keep the evidence and reconciliation record in your private grading files.')
bullet('Do not overwrite another TA\'s entries or the median formulas. Leave unresolved grades blank and identify the pending issue using the instructor\'s tracking method.')
bullet('Check the saved values against your final records and notify the instructor when your assigned entry work is complete. Each TA must enter a grade for every student once resolved.')
p('AI_Assigned_Scores.csv remains the AI-only local tracker. The shared spreadsheet is the class record of reconciled TA grades and calculated median project grades.')

task(9, 'Verify the medians and authorize posting', 'Instructor')
p('Confirm that every expected TA has entered a finalized grade for each student. Verify the formula ranges and roster alignment. A displayed median is not ready for posting if a required TA grade is still missing; a median formula may otherwise ignore blank cells.')
p('The project grade is the median of final TA totals. AI scores do not receive an extra vote, and criterion ratings are not aggregated separately. For an even TA count, the median is the mean of the middle two totals. For example, 82.5, 85.0, and 91.3 yield 85.0.')
p('Review material TA disagreements and unresolved exceptions. Apply any instructor-defined adjustments or display rounding consistently and identify the exact final value to post. Confirm the Canvas assignment, point scale, visibility/release plan, and posting responsibilities. Tell TAs which rows are approved for posting.')

task(10, 'Post the finalized grades to Canvas', 'TA', page_before=True)
p('Post only the instructor-approved final values from the shared spreadsheet. Each TA handles the students or sections assigned for posting; this allocation is separate from the requirement that every TA grades every student.')
for text in [
    'Open the correct Canvas course Gradebook and Project 1 assignment. Match each approved row to the correct student using ID and section.',
    'Enter the approved median project grade in the assignment\'s grade cell. Confirm the assignment point scale matches the instructor\'s instructions. If using a CSV import, use the instructor-approved Canvas-compatible file and review all mapped changes before accepting it.',
    'Verify that Canvas saved the correct value for each student. Resolve any mismatch with the instructor before proceeding with that row. Leave unresolved or unapproved students pending.',
    'Follow the instructor\'s release schedule. Entering a score and making it visible are separate under a manual posting policy: use the assignment\'s Post Grades action when instructed. Under an automatic policy, entering a grade can make it visible immediately.',
    'Confirm the intended grades are visible after release. Record who posted each grade, the date, and completion status in the instructor\'s designated tracking area or a posting log. Preserve the supporting grading records.'
]:
    p(text,'List Number')
p('The grading skill\'s Canvas access is for collecting originals and supporting assessment. Complete this final posting task through the instructor-approved Canvas process; the blind assessment prompt does not authorize the model to post grades.')

doc.add_heading('Records to retain', 2)
table(['Record','Owner and purpose'],[
    ('Original submissions and intake manifests','TA: preserve the files and attempts reviewed.'),
    ('Initial TA assessment and reconciliation record','TA: retain independent judgments, final choices, evidence, and reasons for changes.'),
    ('AI_Assessment.xlsx and AI_Assigned_Scores.csv','TA: retain the AI assessment and verification limits; keep separate from TA final decisions.'),
    ('Shared Google Drive grading spreadsheet','Instructor and TAs: record reconciled TA totals and formula-calculated median grades.'),
    ('Canvas posting record','TA: confirm the approved values were entered and released as instructed.')
],[2.6,4.3])
doc.add_heading('Completion check',2)
p('All roster students are accounted for; required TA grades are finalized; median rows are approved; Canvas values match the shared spreadsheet; release visibility is correct; posting completion and supporting evidence are retained.')
doc.add_heading('Reference links',2)
link('Grading skill folder and supporting resources','https://github.com/cdbale/codex-skills/tree/main/skills/grade-mktg-411-project-1')
link('Canvas guide to entering and editing grades','https://community.instructure.com/en/kb/articles/660852-how-do-i-enter-and-edit-grades-in-the-gradebook')
link('Canvas guide to posting assignment grades','https://community.instructure.com/en/kb/articles/660846-how-do-i-post-grades-for-an-assignment-in-the-gradebook')

doc.save(OUT)
print(OUT.resolve())
print('Words:',len(' '.join(p.text for p in doc.paragraphs).split()))
